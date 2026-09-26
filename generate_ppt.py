import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors (Midnight Cyan & Electric Blue)
    DARK_BG = RGBColor(4, 7, 17)        # #040711
    CARD_BG = RGBColor(11, 20, 42)      # #0B142A
    CYAN_ACCENT = RGBColor(6, 182, 212) # #06B6D4
    SKY_BLUE = RGBColor(56, 189, 248)   # #38BDF8
    WHITE = RGBColor(240, 253, 250)     # #F0FDFA
    TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8
    BORDER_CYAN = RGBColor(8, 145, 178) # #0891B2

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    def add_header(slide, title_text, category_text="SKILLSYNC AI • PROJECT PRESENTATION"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN_ACCENT

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Ambient Card Center
    add_card(s1, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1), bg_color=RGBColor(8, 15, 32), border_color=CYAN_ACCENT)

    tb = s1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.33), Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = "⚡ SkillSync AI"
    p.font.name = "Arial"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    tb2 = s1.shapes.add_textbox(Inches(2.0), Inches(3.0), Inches(9.33), Inches(0.8))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = "Adaptive AI-Powered Education Platform"
    p2.font.name = "Arial"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = CYAN_ACCENT
    p2.alignment = PP_ALIGN.CENTER

    tb3 = s1.shapes.add_textbox(Inches(2.2), Inches(3.8), Inches(8.93), Inches(1.5))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = "Learn what you need. Not what everyone else gets.\nA closed-loop adaptive learning engine that diagnoses conceptual gaps, generates targeted roadmaps, and tutors with real-time feedback."
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Core Problem: One-Size-Fits-All Learning")

    cards_data = [
        ("01", "Static Curriculums", "Every student receives the exact same playlist or textbook chapter, regardless of prior knowledge or specific learning speed."),
        ("02", "Hidden Conceptual Gaps", "Students often don't realize where they are weak until exam day, leading to compounding confusion in foundational topics."),
        ("03", "Passive & Inefficient", "Rewatching full 60-minute video lectures to fix a single 2-minute misunderstanding leads to burnout and high dropouts.")
    ]

    for i, (num, title, desc) in enumerate(cards_data):
        left = Inches(0.9 + i * 3.9)
        add_card(s2, left, Inches(2.1), Inches(3.6), Inches(4.3))

        tb = s2.shapes.add_textbox(left + Inches(0.3), Inches(2.4), Inches(3.0), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = num
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = CYAN_ACCENT

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(10)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(14)

    # ==========================================
    # SLIDE 3: The Solution
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "The Solution: Continuous 5-Step Adaptive Loop")

    loop_steps = [
        ("Step 1", "Diagnostic Assessment", "Targeted 10-minute quiz reveals exact baseline."),
        ("Step 2", "AI Learning Profile", "Subtopic breakdown (Weak, Medium, Strong %) generated."),
        ("Step 3", "Dynamic Study Plan", "Prioritizes high-leverage weak areas first."),
        ("Step 4", "Context-Aware Tutor", "AI explains concepts with direct reference to past mistakes."),
        ("Step 5", "Real-Time Mastery Loop", "Scores recalculate instantly after practice sessions.")
    ]

    for i, (step, title, desc) in enumerate(loop_steps):
        left = Inches(0.9 + i * 2.35)
        add_card(s3, left, Inches(2.2), Inches(2.15), Inches(4.2))

        tb = s3.shapes.add_textbox(left + Inches(0.2), Inches(2.4), Inches(1.75), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_step = tf.paragraphs[0]
        p_step.text = step
        p_step.font.size = Pt(12)
        p_step.font.bold = True
        p_step.font.color.rgb = CYAN_ACCENT

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(12)

    # ==========================================
    # SLIDE 4: Key Platform Features
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Key Features & Capabilities")

    features = [
        ("🎯 Real-Time Diagnostics", "Automated baseline calculation on core subjects (e.g. DBMS, Normalization, Transactions) with instant competency scores."),
        ("🧠 Context-Aware AI Tutor", "Tutor remembers recent diagnostic mistakes and tailors explanation analogies to the student's exact knowledge level."),
        ("📊 Dynamic Mastery Dashboard", "Visual progress breakdown with topic-by-topic status bars, study streaks, and high-priority action items."),
        ("⚡ 1-Click Instant Demo", "Built-in student simulation profile (Alex Rivera) allowing judges and testers to test the adaptive cycle immediately.")
    ]

    for i, (title, desc) in enumerate(features):
        row = i // 2
        col = i % 2
        left = Inches(0.9 + col * 5.9)
        top = Inches(2.1 + row * 2.4)
        add_card(s4, left, top, Inches(5.6), Inches(2.1))

        tb = s4.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), Inches(5.0), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(8)

    # ==========================================
    # SLIDE 5: Tech Stack & Architecture
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Technology Stack & Architecture")

    stack = [
        ("Frontend & UI", "Next.js 16 (App Router + Turbopack)\nReact 19 & TypeScript\nTailwind CSS v4 Modern Glassmorphism\nLucide Icons"),
        ("3D Graphics & Visuals", "Three.js 3D Physics & Crystalline Universe\nReal-time Mouse Cursor Parallax & Lighting\nAdditive Blending Sparkle Particles"),
        ("Backend & Auth", "Next.js Server Actions & API Routes\nNextAuth.js Session & Security\nBcryptjs Password Encryption\nZod Schema Validation"),
        ("Data & Intelligence", "Prisma ORM with SQLite (dev.db)\nGoogle Gemini Generative AI\nAdaptive Scoring Algorithms\nInstant Seeding Pipeline")
    ]

    for i, (title, desc) in enumerate(stack):
        left = Inches(0.9 + i * 2.9)
        add_card(s5, left, Inches(2.1), Inches(2.7), Inches(4.3))

        tb = s5.shapes.add_textbox(left + Inches(0.25), Inches(2.4), Inches(2.2), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN_ACCENT

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = WHITE
        p_d.space_before = Pt(14)

    # ==========================================
    # SLIDE 6: UI/UX Aesthetic Design
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "UI/UX Design: Midnight Cyan Aesthetic")

    design_points = [
        ("🌌 Midnight Space Canvas", "Deep pitch black background (#040711) with neon cyan highlights that minimize eye fatigue during late-night study sessions."),
        ("💎 3D Three.js Crystal Background", "Interactive geometric crystals float, refract light, and react dynamically to mouse gestures, giving a premium tactile feel."),
        ("🛡️ Glassmorphism & Micro-Interactions", "Subtle glowing borders, translucent cards, and animated badges create high visual hierarchy and engagement.")
    ]

    for i, (title, desc) in enumerate(design_points):
        top = Inches(2.1 + i * 1.5)
        add_card(s6, Inches(0.9), top, Inches(11.5), Inches(1.25))

        tb = s6.shapes.add_textbox(Inches(1.2), top + Inches(0.2), Inches(10.9), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN_ACCENT

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = WHITE
        p_d.space_before = Pt(4)

    # ==========================================
    # SLIDE 7: Future Roadmap
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Future Roadmap & Expansion")

    roadmap = [
        ("Phase 1: Multi-Subject", "Scale beyond DBMS to Operating Systems, Computer Networks, and Data Structures & Algorithms."),
        ("Phase 2: Teacher Dashboard", "Analytics portal for educators to identify class-wide misconceptions and weak topic clusters."),
        ("Phase 3: Peer Study Rooms", "Collaborative problem-solving with AI moderator facilitating group peer reviews.")
    ]

    for i, (title, desc) in enumerate(roadmap):
        left = Inches(0.9 + i * 3.9)
        add_card(s7, left, Inches(2.1), Inches(3.6), Inches(4.3))

        tb = s7.shapes.add_textbox(left + Inches(0.3), Inches(2.4), Inches(3.0), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN_ACCENT

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(14)

    # ==========================================
    # SLIDE 8: Conclusion & Demo
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)

    add_card(s8, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1), bg_color=RGBColor(8, 15, 32), border_color=CYAN_ACCENT)

    tb = s8.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.33), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = "Arial"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    tb2 = s8.shapes.add_textbox(Inches(2.0), Inches(3.2), Inches(9.33), Inches(0.8))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = "SkillSync AI — Empowering Every Student with Adaptive Intelligence"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = CYAN_ACCENT
    p2.alignment = PP_ALIGN.CENTER

    tb3 = s8.shapes.add_textbox(Inches(2.2), Inches(4.2), Inches(8.93), Inches(1.2))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = "🌐 Live Demo: http://localhost:3000\n💻 Tech: Next.js 16 • Three.js • Prisma • Gemini AI\n✨ Ready for Questions & Interactive Walkthrough"
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = WHITE
    p3.alignment = PP_ALIGN.CENTER

    # Save
    output_path = r"c:\Users\admin\OneDrive\Desktop\SkillSync-main\SkillSync-main\SkillSync_AI_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
